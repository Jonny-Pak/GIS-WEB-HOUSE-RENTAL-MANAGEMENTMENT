import io
from datetime import datetime, date

import openpyxl
from django.contrib.auth.decorators import user_passes_test
from django.http import HttpResponse
from django.shortcuts import render
from django.utils.timezone import make_aware, is_naive

from custom_admin.views.helpers import _is_admin_user
from houses.models import House


@user_passes_test(_is_admin_user, login_url='custom_admin_login')
def export_available_houses_excel(request):
    """Xuất Excel danh sách nhà chưa cho thuê (status=available)."""
    houses = House.objects.filter(status='available').select_related('owner').order_by('-created_at')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Nhà chưa cho thuê"

    # Header
    headers = [
        'STT', 'Tên nhà / Tiêu đề', 'Chủ nhà', 'SĐT liên hệ',
        'Loại hình', 'Địa chỉ', 'Diện tích (m²)',
        'Giá thuê (VNĐ/tháng)', 'Tiền cọc (VNĐ)', 'Ngày đăng',
    ]
    ws.append(headers)

    # Căn chỉnh độ rộng cột
    col_widths = [5, 40, 20, 15, 20, 40, 15, 20, 20, 15]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    # Dữ liệu
    for idx, house in enumerate(houses, start=1):
        price_display = f"{house.price:,}" if house.price else "Thỏa thuận"
        deposit_display = f"{house.deposit:,}" if house.deposit else "0"
        owner_name = house.owner.get_full_name() or house.owner.username if house.owner else "-"
        created_at = house.created_at.strftime("%d/%m/%Y") if house.created_at else "-"

        ws.append([
            idx,
            house.name,
            owner_name,
            house.owner_phone or "-",
            house.get_house_type_display(),
            house.address,
            house.area,
            price_display,
            deposit_display,
            created_at,
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    today_str = date.today().strftime("%d-%m-%Y")
    filename = f"nha_chua_cho_thue_{today_str}.xlsx"

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@user_passes_test(_is_admin_user, login_url='custom_admin_login')
def dashboard_filter_listings(request):
    """Lọc danh sách tin đã được duyệt theo khoảng ngày và trả về kết quả + xuất Excel."""
    date_from_str = request.GET.get('date_from', '').strip()
    date_to_str = request.GET.get('date_to', '').strip()
    export = request.GET.get('export', '')

    houses = None
    error = None
    date_from = None
    date_to = None

    if date_from_str and date_to_str:
        try:
            date_from = datetime.strptime(date_from_str, "%Y-%m-%d").date()
            date_to = datetime.strptime(date_to_str, "%Y-%m-%d").date()

            if date_from > date_to:
                error = "Ngày bắt đầu không được lớn hơn ngày kết thúc."
            else:
                # Lấy tin đã duyệt (available hoặc rented) được đăng trong khoảng ngày
                houses = House.objects.filter(
                    status__in=['available', 'rented'],
                    created_at__date__gte=date_from,
                    created_at__date__lte=date_to,
                ).select_related('owner').order_by('-created_at')

                # Nếu yêu cầu xuất Excel
                if export == '1':
                    return _export_filtered_excel(houses, date_from_str, date_to_str)

        except ValueError:
            error = "Định dạng ngày không hợp lệ. Vui lòng chọn lại."

    context = {
        'houses': houses,
        'date_from': date_from_str,
        'date_to': date_to_str,
        'error': error,
        'total': houses.count() if houses is not None else 0,
    }
    return render(request, 'custom_admin/filter_listings.html', context)


def _export_filtered_excel(houses, date_from_str, date_to_str):
    """Tạo file Excel cho danh sách tin lọc theo ngày."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tin đăng theo ngày"

    # Tiêu đề báo cáo
    try:
        d_from = datetime.strptime(date_from_str, "%Y-%m-%d").strftime("%d/%m/%Y")
        d_to = datetime.strptime(date_to_str, "%Y-%m-%d").strftime("%d/%m/%Y")
    except Exception:
        d_from = date_from_str
        d_to = date_to_str

    ws.append([f"Danh sách tin đăng đã duyệt từ {d_from} đến {d_to}"])
    ws.append([f"Tổng số tin: {houses.count()}"])
    ws.append([])  # dòng trống

    headers = [
        'STT', 'Tên nhà / Tiêu đề', 'Chủ nhà', 'SĐT liên hệ',
        'Loại hình', 'Địa chỉ', 'Diện tích (m²)',
        'Giá thuê (VNĐ/tháng)', 'Trạng thái', 'Ngày đăng',
    ]
    ws.append(headers)

    col_widths = [5, 40, 20, 15, 20, 40, 15, 20, 15, 15]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    for idx, house in enumerate(houses, start=1):
        price_display = f"{house.price:,}" if house.price else "Thỏa thuận"
        owner_name = house.owner.get_full_name() or house.owner.username if house.owner else "-"
        created_at = house.created_at.strftime("%d/%m/%Y") if house.created_at else "-"

        ws.append([
            idx,
            house.name,
            owner_name,
            house.owner_phone or "-",
            house.get_house_type_display(),
            house.address,
            house.area,
            price_display,
            house.get_status_display(),
            created_at,
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"tin_dang_{date_from_str}_den_{date_to_str}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
